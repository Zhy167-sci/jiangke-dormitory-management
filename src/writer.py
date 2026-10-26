"""
writer.py —— 输出层：把统计结果写成带格式的 Excel。

学习要点
--------
1. 数据 和 样式 要分开：aggregate 只产出“数字和文字”，
   这里才决定它们长什么样。将来领导要换版式，只改这个文件。
2. openpyxl 的关键对象：
     Font / PatternFill（字体、底色）
     Border + Side（边框）
     Alignment（对齐、自动换行）
     column_dimensions（列宽）
     page_setup（A4、缩放、页边距）
     freeze_panes（冻结窗格，滚动时表头不动）
3. 数字必须写成数字（float），不要写成字符串 '95'，
   否则在 Excel 里没法排序和求和。
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from aggregate import format_bed
from config import OTHER_BED_TEXT

# ---------------------------------------------------------------- 样式常量
COLOR_HEAD = "FF1F4E79"      # 表头深蓝
COLOR_BAND = "FFF4F7FA"      # 斑马纹浅灰蓝
COLOR_LINE = "FFBFC7D2"      # 边框灰
COLOR_WARN = "FFB45309"      # 空床提示橙
COLOR_BAD = "FFFCE8E8"       # 未查到浅红

FONT_NAME = "微软雅黑"

_thin = Side(style="thin", color=COLOR_LINE)
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _font(size=10, bold=False, color=None):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


def _fill(color):
    return PatternFill(fill_type="solid", fgColor=color)


# ---------------------------------------------------------------- 通用写表
def write_table(ws, title, subtitle, headers, body, widths,
                landscape=True, bad_rows=None):
    """往工作表里写“标题 + 副标题 + 表头 + 数据”。

    参数
    ----
    body     : 二维列表，每个元素是一行（已渲染好的文字或数字）
    widths   : 每列宽度列表
    bad_rows : 需要标红的行号集合（1 基，相对 body）
    """
    ncol = len(headers)

    # 标题（合并整行、居中、加粗）
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = _font(15, True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # 副标题
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    cell = ws.cell(row=2, column=1, value=subtitle)
    cell.font = _font(9, False, "FF6B7280")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # 表头
    for c, text in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c, value=text)
        cell.font = _font(10, True, "FFFFFFFF")
        cell.fill = _fill(COLOR_HEAD)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[3].height = 32

    # 数据
    for r, row in enumerate(body, start=4):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = _font(10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER

            if (bad_rows and (r - 3) in bad_rows) or value == "未查到":
                cell.fill = _fill(COLOR_BAD)
            elif r % 2 == 0:
                cell.fill = _fill(COLOR_BAND)

            if value == "空床":
                cell.font = _font(10, True, COLOR_WARN)

        ws.row_dimensions[r].height = 23

    # 列宽
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 冻结前 3 行（标题两行 + 表头），滚动时表头不跑
    ws.freeze_panes = "A4"

    # 打印设置：A4、横向/纵向、缩放到一页宽
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


# ---------------------------------------------------------------- 按班级
def save_class_sheet(path, cls, week, rows):
    """一个班一个文件。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"

    headers = ["楼栋", "寝室号", "床位1", "床位2", "床位3", "床位4", "个人均分", "寝室得分"]
    body = []
    bad = set()
    for i, r in enumerate(rows, start=1):
        # 按班级表：不是本班的床位显示「—」，不是「空床」
        beds = [format_bed(b, OTHER_BED_TEXT) for b in r["床位"]]
        if any(b.endswith("未查到") for b in beds) or r["寝室得分"] is None:
            bad.add(i)
        body.append([
            r["楼栋"], r["寝室号"], *beds,
            r["个人均分"] if r["个人均分"] is not None else "—",
            r["寝室得分"] if r["寝室得分"] is not None else "未查到",
        ])

    subtitle = (f"{week}　寝室 {len(rows)} 间 / {sum(1 for r in rows for b in r['床位'] if b)} 人"
                f"　寝室均分 {_avg_of([r['寝室得分'] for r in rows])}")

    write_table(ws, f"{cls}　查寝得分汇总", subtitle, headers, body,
                widths=[7, 10, 17, 17, 17, 17, 11, 11],
                landscape=False, bad_rows=bad)

    wb.save(str(path))
    return path


# ---------------------------------------------------------------- 按楼栋
def save_building_sheet(path, building, week, rows):
    """一栋楼一个文件。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"

    headers = ["寝室号", "班级", "床位1", "床位2", "床位3", "床位4", "个人均分", "寝室得分"]
    body = []
    bad = set()
    for i, r in enumerate(rows, start=1):
        beds = [format_bed(b) for b in r["床位"]]
        if r["寝室得分"] is None:
            bad.add(i)
        body.append([
            r["寝室号"], r["班级"], *beds,
            r["个人均分"] if r["个人均分"] is not None else "—",
            r["寝室得分"] if r["寝室得分"] is not None else "未查到",
        ])

    subtitle = (f"{week}　寝室 {len(rows)} 间"
                f"　寝室均分 {_avg_of([r['寝室得分'] for r in rows])}")

    write_table(ws, f"{building} 栋 查寝得分汇总", subtitle, headers, body,
                widths=[10, 30, 17, 17, 17, 17, 11, 11],
                landscape=True, bad_rows=bad)

    wb.save(str(path))
    return path


def _avg_of(values):
    vals = [v for v in values if v is not None]
    return round(sum(vals) / len(vals), 2) if vals else "—"
